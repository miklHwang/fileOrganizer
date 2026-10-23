from __future__ import annotations

import io
import os
import time
from datetime import datetime

import pytest
from PIL import Image

from fixtures.fake_tree import _lp, build_fake_tree, snapshot


def test_build_is_fast(tmp_path):
    t0 = time.time()
    build_fake_tree(tmp_path)
    assert time.time() - t0 < 30


def test_every_handle_exists(fake_tree):
    assert len(fake_tree.files) > 30
    for name, p in fake_tree.files.items():
        assert os.path.exists(_lp(p)), name
    for lst in (fake_tree.cloud_only, fake_tree.hidden, fake_tree.protected,
                fake_tree.excluded, fake_tree.atomic_dirs):
        assert lst
        for p in lst:
            assert os.path.exists(_lp(p)), p
    assert len(fake_tree.atomic_dirs) == 3
    assert len(fake_tree.cloud_only) == 3
    for p in fake_tree.folders.values():
        assert p.is_dir()
    assert fake_tree.folders["screenshots"] == fake_tree.folders["pictures"] / "Screenshots"


def test_png_dimensions(fake_tree):
    exp = {"screenshot_desktop": (1920, 1080), "screenshot_capture": (2560, 1440),
           "screenshot_phone": (1080, 2400), "existing_review_screenshot": (1920, 1080),
           "drawing_png": (500, 500)}
    for k, size in exp.items():
        with Image.open(fake_tree.files[k]) as im:
            assert im.size == size, k


def test_exif(fake_tree):
    with Image.open(fake_tree.files["photo_exif"]) as im:
        ex = im.getexif()
        assert ex[0x010F] == "Canon" and ex[0x0110] == "EOS 80D"
        assert ex.get_ifd(0x8769)[0x9003] == "2024:07:14 10:10:10"
    with Image.open(fake_tree.files["photo_exif2"]) as im:
        ex = im.getexif()
        assert ex[0x010F] == "NIKON"
        assert ex.get_ifd(0x8769)[0x9003] == "2023:01:02 09:00:00"
    with Image.open(fake_tree.files["photo_near"]) as im:
        assert im.size == (576, 432)


def test_pdfs(fake_tree):
    from pypdf import PdfReader

    def text(k):
        r = PdfReader(str(fake_tree.files[k]))
        return "".join(pg.extract_text() or "" for pg in r.pages)

    assert "Invoice number" in text("invoice_pdf")
    assert "Closing balance" in text("statement_pdf")
    assert "Syllabus" in text("syllabus_pdf")
    assert text("scanned_pdf").strip() == ""


def test_docx(fake_tree):
    import docx

    d = docx.Document(str(fake_tree.files["agenda_docx"]))
    assert "Project kickoff meeting" in [p.text for p in d.paragraphs]


def test_duplicate_groups_identical(fake_tree):
    assert len(fake_tree.duplicate_groups) == 2
    for g in fake_tree.duplicate_groups:
        assert len({p.read_bytes() for p in g}) == 1
    a, b = fake_tree.collision_pair
    assert a.name == b.name and a.read_bytes() != b.read_bytes()


def test_version_chain_and_mtimes(fake_tree):
    assert [p.name for p in fake_tree.version_chain] == [
        "report v1.docx", "report v2.docx", "report final_FINAL.docx"]
    assert [datetime.fromtimestamp(p.stat().st_mtime).day for p in fake_tree.version_chain] == [1, 2, 3]
    now = datetime(2025, 6, 1, 12)
    t = build_fake_tree(fake_tree.root / "other", now=now)
    age = now - datetime.fromtimestamp(t.files["installer_old"].stat().st_mtime)
    assert age.days == 200
    assert (now - datetime.fromtimestamp(t.files["installer_new"].stat().st_mtime)).days == 5
    assert datetime.fromtimestamp(t.files["misc"].stat().st_mtime) == datetime(2024, 1, 10, 12)


def test_long_path(fake_tree):
    rel = str(fake_tree.long_path_file.relative_to(fake_tree.root))
    assert len(str(fake_tree.long_path_file)) > 260 or len(rel) > 240
    assert os.path.exists(_lp(fake_tree.long_path_file))
    assert len(str(fake_tree.long_path_file.relative_to(fake_tree.folders["documents"]))) > 240


def test_config_overrides_and_all_files(fake_tree):
    o = fake_tree.config_overrides()
    assert o["work_dir"] == str(fake_tree.root / "tidyup_data")
    assert o["destination"] == {"root": None}
    assert o["folders"]["home"] == str(fake_tree.folders["home"])
    files = fake_tree.all_files()
    assert fake_tree.long_path_file in files
    assert fake_tree.files["invoice_pdf"] in files


def test_snapshot(fake_tree):
    home = fake_tree.folders["home"]
    s1 = snapshot(home)
    assert s1 == snapshot(home)
    assert s1["Documents"] is None
    assert "Documents/.hidden_notes.txt" in s1
    f = fake_tree.files["misc"]
    st = f.stat()

    data = f.read_bytes()
    f.write_bytes(data[:-1] + bytes([data[-1] ^ 1]))
    os.utime(f, (st.st_atime, st.st_mtime))
    s2 = snapshot(home)
    assert s2 != s1
    f.write_bytes(data)
    os.utime(f, (st.st_atime, st.st_mtime))
    assert snapshot(home) == s1

    os.utime(f, (st.st_atime, st.st_mtime + 100))
    assert snapshot(home) != s1
    os.utime(f, (st.st_atime, st.st_mtime))

    f.rename(f.with_name("renamed.bin"))
    assert snapshot(home) != s1


def test_cfg_and_spy_fixtures_exist(request):
    # these need tidyup.winutil / tidyup.config (built in parallel tasks)
    for n in ("cfg", "mock_attrs", "cross_volume", "open_spy"):
        assert request._fixturemanager.getfixturedefs(n, request.node)


def test_open_spy_before_fake_tree(open_spy, fake_tree):
    from tidyup import winutil

    assert fake_tree.files
    p = fake_tree.files["misc"]
    with open(p, "rb") as fh:
        fh.read(1)
    assert winutil.norm_key(str(p)) in open_spy


def test_open_spy_records_image_and_pdf(open_spy, fake_tree):
    from pypdf import PdfReader
    from tidyup import winutil

    img = fake_tree.files["screenshot_desktop"]
    pdf = fake_tree.files["invoice_pdf"]
    with Image.open(img):
        pass
    PdfReader(str(pdf))
    assert winutil.norm_key(str(img)) in open_spy
    assert winutil.norm_key(str(pdf)) in open_spy


def test_open_spy_pdfreader_stays_a_class(open_spy, fake_tree):
    import pypdf
    from tidyup import winutil

    pdf = fake_tree.files["invoice_pdf"]
    r = pypdf.PdfReader(str(pdf))
    assert isinstance(r, pypdf.PdfReader)
    assert isinstance(pypdf.PdfReader, type)
    assert winutil.norm_key(str(pdf)) in open_spy


def test_open_spy_pdfreader_subclass(open_spy, fake_tree):
    import pypdf
    from tidyup import winutil

    class Sub(pypdf.PdfReader):
        pass

    class Sub2(pypdf.PdfReader):
        def __init__(self, path):
            super().__init__(path)

    p1, p2 = fake_tree.files["invoice_pdf"], fake_tree.files["statement_pdf"]
    assert isinstance(Sub(str(p1)), pypdf.PdfReader)
    assert isinstance(Sub2(str(p2)), pypdf.PdfReader)
    assert winutil.norm_key(str(p1)) in open_spy
    assert winutil.norm_key(str(p2)) in open_spy


def test_open_spy_pdfreader_bound_before_fixture(request, fake_tree):
    from pypdf import PdfReader  # bound before open_spy is active
    from tidyup import winutil

    request.getfixturevalue("open_spy")
    opened = request.getfixturevalue("open_spy")
    pdf = fake_tree.files["syllabus_pdf"]
    PdfReader(str(pdf))
    assert winutil.norm_key(str(pdf)) in opened


def test_open_spy_pdfreader_file_object_records_name(open_spy, fake_tree):
    import pypdf
    from tidyup import winutil

    pdf = fake_tree.files["invoice_pdf"]
    with open(pdf, "rb") as fh:
        pypdf.PdfReader(fh)
    assert winutil.norm_key(str(pdf)) in open_spy


def test_open_spy_docx_document(open_spy, fake_tree):
    import inspect

    import docx
    from tidyup import winutil

    path = fake_tree.files["agenda_docx"]
    d = docx.Document(str(path))
    if inspect.isclass(docx.Document):
        assert isinstance(d, docx.Document)
    assert d.paragraphs
    assert winutil.norm_key(str(path)) in open_spy


def test_open_spy_pil_open_recorded(open_spy, fake_tree):
    from tidyup import winutil

    img = fake_tree.files["photo_exif"]
    with Image.open(img) as im:
        assert im.size
    assert winutil.norm_key(str(img)) in open_spy
