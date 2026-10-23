"""Content readers. Every reader refuses cloud-only files (never opens them) and
returns None / "" for corrupt or unreadable files."""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from tidyup import winutil

IMAGE_EXTS = frozenset({".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tif", ".tiff"})


@dataclass(frozen=True)
class ImageInfo:
    width: int
    height: int
    camera_make: str | None
    camera_model: str | None
    exif_datetime: datetime | None
    phash: str | None


def _blocked(path: Path, is_cloud_only: bool) -> bool:
    """True when the file must not be opened (flag or live attribute check)."""
    return bool(is_cloud_only) or winutil.is_cloud_only(path)


def _clean(v: object) -> str | None:
    if isinstance(v, bytes):
        v = v.decode("utf-8", "ignore")
    if not isinstance(v, str):
        return None
    v = v.replace("\x00", "").strip()
    return v or None


def _parse_exif_dt(v: object) -> datetime | None:
    v = _clean(v)
    if not v:
        return None
    for fmt in ("%Y:%m:%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(v[:19], fmt)
        except ValueError:
            continue
    return None


def read_image_info(path: Path, *, is_cloud_only: bool, want_phash: bool = True) -> ImageInfo | None:
    if _blocked(path, is_cloud_only):
        return None
    try:
        from PIL import Image

        with Image.open(winutil.long_path(path)) as img:
            width, height = img.size
            make = model = None
            dt = None
            try:
                exif = img.getexif()
                make = _clean(exif.get(0x010F))
                model = _clean(exif.get(0x0110))
                try:
                    dt = _parse_exif_dt(exif.get_ifd(0x8769).get(36867))
                except Exception:
                    dt = None
                if dt is None:
                    dt = _parse_exif_dt(exif.get(306))
            except Exception:
                pass
            ph = None
            if want_phash:
                try:
                    import imagehash

                    img.load()
                    ph = str(imagehash.phash(img))
                except Exception:
                    ph = None
            return ImageInfo(width, height, make, model, dt, ph)
    except Exception:
        return None


def pdf_text(path: Path, *, is_cloud_only: bool, max_pages: int = 2, max_chars: int = 20000) -> str:
    if _blocked(path, is_cloud_only):
        return ""
    try:
        import pypdf

        reader = pypdf.PdfReader(winutil.long_path(path))
        parts: list[str] = []
        for page in list(reader.pages)[:max_pages]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(parts).strip()[:max_chars]
    except Exception:
        return ""


def docx_text(path: Path, *, is_cloud_only: bool, max_chars: int = 20000) -> str:
    if _blocked(path, is_cloud_only):
        return ""
    try:
        import docx

        doc = docx.Document(winutil.long_path(path))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts).strip()[:max_chars]
    except Exception:
        return ""


def thumbnail_data_uri(path: Path, *, is_cloud_only: bool, max_px: int = 160) -> str | None:
    if _blocked(path, is_cloud_only):
        return None
    try:
        from PIL import Image, ImageOps

        with Image.open(winutil.long_path(path)) as img:
            img.load()
            try:
                img = ImageOps.exif_transpose(img)
            except Exception:
                pass
            rgba = img.convert("RGBA")
            bg = Image.new("RGB", rgba.size, (255, 255, 255))
            bg.paste(rgba, mask=rgba.split()[3])
            bg.thumbnail((max_px, max_px))
            buf = io.BytesIO()
            bg.save(buf, "JPEG", quality=70)
        return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    except Exception:
        return None
